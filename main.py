"""
Автоматизация отчёта по КРГ (инвентаризация).

Приложение выполняет три шага, которые ранее делались вручную:
    1. Строит сводную по товарным группам на основе листа «Данные КРГ»
       (суммы Дельта, СуммаРозница, СуммаПриход).
    2. Записывает сводные значения на лист «Отчет по ТГ» (столбцы T/U/V).
    3. На листе «Готовый отчет по ТГ» записывает те же значения в T/U/V
       и добавляет корректировку в столбцы I/J/K по формуле:
           =<базовое значение>+Таблица[[#This Row],[КРГ, ...]]
       Базовое значение берётся с листа «Отчет по ТГ» (эталонного,
       не изменяемого вручную), поэтому расчёт идемпотентен —
       повторная обработка не приводит к двойному учёту корректировки.
"""

from __future__ import annotations

from collections import defaultdict
from copy import copy
from dataclasses import dataclass, field
from io import BytesIO
from typing import BinaryIO

import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

# Константы

SHEET_SOURCE = "Данные КРГ"
SHEET_REPORT = "Отчет по ТГ"
SHEET_READY = "Готовый отчет по ТГ"

COL_SOURCE_TG = 2      # B — ТоварнаяГруппа
COL_SOURCE_DELTA = 12  # L — Дельта
COL_SOURCE_RETAIL = 13  # M — СуммаРозница
COL_SOURCE_INCOME = 14  # N — СуммаПриход

HEADERS_REPORT_TUV = (
    "Дельта КРГ, шт",
    "Дельта КРГ в средневзвеш.розн. ценах, руб.",
    "Дельта КРГ в приходных ценах, руб.",
)
HEADERS_REPORT_IJK = (
    "Фактический остаток товара на момент сплошной инвентаризации, шт.",
    "Фактический остаток в средневзвеш.розн. ценах, руб.",
    "Фактический остаток в приходных ценах, руб.",
)
HEADERS_READY_TUV = (
    "КРГ, шт",
    "КРГ, средневзвешенная, руб.",
    "КРГ, приходная, руб.",
)
HEADERS_READY_IJK = (
    "Фактический остаток товара на момент сплошной инвентаризации, шт. + КРГ",
    "Фактический остаток в средневзвеш.розн. ценах, руб. + КРГ",
    "Фактический остаток в приходных ценах, руб. + КРГ",
)

FILL_ADJUSTED = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")

ColTriple = tuple[int, int, int]
ValueTriple = tuple[object, object, object]


# Вспомогательные функции

def to_number(value: object) -> float:
    """Приводит значение ячейки Excel к числу с плавающей точкой."""
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(" ", "").replace("\xa0", "").replace(",", "."))
    except ValueError:
        return 0.0


def normalize(value: object) -> str:
    """Нормализует название товарной группы для сопоставления между листами."""
    if value is None:
        return ""
    return " ".join(str(value).replace("\xa0", " ").split()).lower()


def find_header_columns(ws: Worksheet, headers: tuple[str, ...], row: int = 1) -> ColTriple:
    """Возвращает номера столбцов по точным названиям заголовков."""
    header_row = {cell.value: cell.column for cell in ws[row]}
    missing = [h for h in headers if h not in header_row]
    if missing:
        raise ValueError(f'На листе "{ws.title}" не найдены столбцы: {", ".join(missing)}')
    col_t, col_u, col_v = (header_row[h] for h in headers)
    return col_t, col_u, col_v


def find_table_name(ws: Worksheet, default: str = "Таблица5") -> str:
    """Возвращает имя первой Excel-таблицы на листе (для ссылок в формулах)."""
    return next(iter(ws.tables), default)


def build_tg_row_index(ws: Worksheet, tg_column: int = 3) -> dict[str, int]:
    """Сопоставляет нормализованное название ТГ с номером строки на листе."""
    index: dict[str, int] = {}
    for row in range(2, ws.max_row + 1):
        tg = ws.cell(row=row, column=tg_column).value
        if tg:
            index[normalize(tg)] = row
    return index


# Основная логика

@dataclass
class TGSummary:
    name: str
    delta: float = 0.0
    retail: float = 0.0
    income: float = 0.0


@dataclass
class ProcessingResult:
    workbook_bytes: BytesIO
    summary: pd.DataFrame
    missing_in_report: list[str] = field(default_factory=list)
    missing_in_ready: list[str] = field(default_factory=list)


def write_tuv(ws: Worksheet, row: int, columns: ColTriple, entry: TGSummary) -> None:
    """Записывает дельту, розницу и приход в три столбца одной строки."""
    col_t, col_u, col_v = columns
    ws.cell(row=row, column=col_t).value = entry.delta
    ws.cell(row=row, column=col_u).value = entry.retail
    ws.cell(row=row, column=col_v).value = entry.income


def build_pivot(ws_source: Worksheet) -> dict[str, TGSummary]:
    """Строит сводную по товарным группам из листа «Данные КРГ»."""
    pivot: dict[str, TGSummary] = defaultdict(lambda: TGSummary(name=""))
    for row in ws_source.iter_rows(min_row=2):
        tg = row[COL_SOURCE_TG - 1].value
        if not tg:
            continue
        entry = pivot[normalize(tg)]
        entry.name = entry.name or str(tg).strip()
        entry.delta += to_number(row[COL_SOURCE_DELTA - 1].value)
        entry.retail += to_number(row[COL_SOURCE_RETAIL - 1].value)
        entry.income += to_number(row[COL_SOURCE_INCOME - 1].value)
    return pivot


def apply_to_report(
    ws_report: Worksheet,
    pivot: dict[str, TGSummary],
) -> tuple[dict[str, ValueTriple], list[str]]:
    """Записывает T/U/V на лист «Отчет по ТГ». Возвращает базовые I/J/K по ТГ (для шага 3)."""
    cols_tuv = find_header_columns(ws_report, HEADERS_REPORT_TUV)
    col_i, col_j, col_k = find_header_columns(ws_report, HEADERS_REPORT_IJK)
    row_index = build_tg_row_index(ws_report)

    base_values: dict[str, ValueTriple] = {}
    missing: list[str] = []

    for key, entry in pivot.items():
        row = row_index.get(key)
        if row is None:
            missing.append(entry.name)
            continue
        write_tuv(ws_report, row, cols_tuv, entry)
        base_values[key] = (
            ws_report.cell(row=row, column=col_i).value,
            ws_report.cell(row=row, column=col_j).value,
            ws_report.cell(row=row, column=col_k).value,
        )

    return base_values, missing


def apply_to_ready(
    ws_ready: Worksheet,
    pivot: dict[str, TGSummary],
    base_values: dict[str, ValueTriple],
) -> list[str]:
    """Записывает T/U/V и формулы I/J/K на лист «Готовый отчет по ТГ»."""
    cols_tuv = find_header_columns(ws_ready, HEADERS_READY_TUV)
    cols_ijk = find_header_columns(ws_ready, HEADERS_READY_IJK)
    row_index = build_tg_row_index(ws_ready)
    table_name = find_table_name(ws_ready)

    missing: list[str] = []

    for key, entry in pivot.items():
        row = row_index.get(key)
        if row is None:
            missing.append(entry.name)
            continue

        write_tuv(ws_ready, row, cols_tuv, entry)

        bases = base_values.get(key, (0, 0, 0))
        for col, header, base in zip(cols_ijk, HEADERS_READY_TUV, bases):
            cell = ws_ready.cell(row=row, column=col)
            cell.value = f"={base}+{table_name}[[#This Row],[{header}]]"
            cell.fill = copy(FILL_ADJUSTED)

    return missing


def process_workbook(uploaded_file: BinaryIO) -> ProcessingResult:
    """Полный цикл обработки загруженного файла."""
    workbook = openpyxl.load_workbook(BytesIO(uploaded_file.getvalue()), data_only=False)

    missing_sheets = [name for name in (SHEET_SOURCE, SHEET_REPORT, SHEET_READY) if name not in workbook.sheetnames]
    if missing_sheets:
        raise ValueError(f'Не найден обязательный лист «{missing_sheets[0]}».')

    pivot = build_pivot(workbook[SHEET_SOURCE])
    base_values, missing_report = apply_to_report(workbook[SHEET_REPORT], pivot)
    missing_ready = apply_to_ready(workbook[SHEET_READY], pivot, base_values)

    workbook.calculation.fullCalcOnLoad = True

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    summary = pd.DataFrame(
        [
            {
                "Товарная группа": e.name,
                "Дельта": e.delta,
                "СуммаРозница": e.retail,
                "СуммаПриход": e.income,
            }
            for e in pivot.values()
        ]
    )

    return ProcessingResult(
        workbook_bytes=output,
        summary=summary,
        missing_in_report=missing_report,
        missing_in_ready=missing_ready,
    )


# Интерфейс

def render_result(result: ProcessingResult) -> None:
    st.success("Обработка завершена успешно.")

    st.subheader("Сводная по товарным группам")
    st.dataframe(result.summary, use_container_width=True, hide_index=True)

    if result.missing_in_report:
        st.warning(
            "Не найдены на листе «Отчет по ТГ»: " + ", ".join(result.missing_in_report)
        )
    if result.missing_in_ready:
        st.warning(
            "Не найдены на листе «Готовый отчет по ТГ»: " + ", ".join(result.missing_in_ready)
        )

    st.download_button(
        label="Скачать обработанный файл",
        data=result.workbook_bytes,
        file_name="Готовый_отчет_КРГ.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )


def main() -> None:
    st.set_page_config(page_title="Автоматизация отчёта КРГ", page_icon="📋", layout="centered")

    st.title("Автоматизация отчёта по КРГ")
    st.caption(
        "Загрузите файл инвентаризации с листами «Данные КРГ», «Отчет по ТГ» "
        "и «Готовый отчет по ТГ». Приложение рассчитает сводную по товарным "
        "группам и внесёт корректировки автоматически."
    )

    uploaded_file = st.file_uploader("Исходный файл (.xlsx)", type=["xlsx"])
    if uploaded_file is None:
        return

    st.info(f"Файл загружен: **{uploaded_file.name}**")
    if not st.button("Обработать файл", type="primary"):
        return

    try:
        with st.spinner("Выполняется обработка..."):
            result = process_workbook(uploaded_file)
        render_result(result)
    except ValueError as exc:
        st.error(str(exc))
    except Exception:
        st.error(
            "Не удалось обработать файл. Проверьте, что структура листов "
            "соответствует шаблону, и попробуйте снова."
        )

main()
